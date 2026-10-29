#!/usr/bin/env python3
"""Independent oracle for the docgen fidelity evidence.

This file deliberately does NOT import the TypeScript generator. It reads the produced
artifacts and the eCap reference artifacts with python-docx 1.2.0 / openpyxl 3.1.5, then
reports structural agreements and differences.

usage:
  oracle.py <ours.xlsx> <ecap.xlsx> <ours.docx> <ecap.docx> \
            <run1.json> <run2.json> <verify.json> \
            <verify-tampered-xlsx.json> <verify-tampered-docx.json> <outdir>

writes <outdir>/oracle.json and <outdir>/result.json; prints a compact summary to stdout.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
import sys
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.oxml.ns import qn


# ---------------------------------------------------------------------------- helpers
def sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    digest.update(Path(path).read_bytes())
    return digest.hexdigest()


TITLE_PREFIX = re.compile(r"[（(]ECAP_([A-Za-z_][A-Za-z0-9_]*)[）)]")
SOURCE_MANUAL = "/home/haroldzhao/eCap/docs/user-manual.md"


def norm_reg(value: Any) -> Any:
    """Strip the known eCap core-register prefix; all other text is untouched.

    Two shapes occur: a bare name (`ECAP_TSCTR`, an overview cell) and a name inside the
    register-chapter heading parentheses (`时间戳计数器（ECAP_TSCTR）`).
    """
    if isinstance(value, str):
        if value.startswith("ECAP_"):
            return value[len("ECAP_"):]
        return TITLE_PREFIX.sub(lambda match: "（" + match.group(1) + "）", value)
    return value


def canon(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def digest(value: Any) -> str:
    return hashlib.sha256(canon(value).encode("utf-8")).hexdigest()


def compact_list(values: list[Any]) -> dict[str, Any]:
    return {"count": len(values), "digest": digest(values)}


def summary(items: list[Any]) -> dict[str, Any]:
    return {"count": len(items), "digest": digest(items)}


# ---------------------------------------------------------------------------- xlsx oracle
def scalar(value: Any) -> Any:
    """Coerce openpyxl descriptor wrappers (e.g. Integer) to plain JSON scalars."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    try:
        return int(value)
    except (TypeError, ValueError):
        return str(value)


def color_of(color: Any) -> Any:
    if color is None:
        return None
    return {"rgb": scalar(getattr(color, "rgb", None)), "theme": scalar(getattr(color, "theme", None)),
            "indexed": scalar(getattr(color, "indexed", None)), "type": scalar(getattr(color, "type", None))}


def side_of(side: Any) -> Any:
    if side is None:
        return None
    return {"style": scalar(side.style), "color": color_of(side.color)}


def xlsx_cell_sig(cell: Any) -> dict[str, Any]:
    font = cell.font
    fill = None
    if getattr(cell.fill, "fill_type", None) not in (None, "none"):
        fill = color_of(cell.fill.fgColor)
    border = cell.border
    return {
        "bold": bool(font.bold),
        "font": font.name,
        "size": float(font.sz) if font.sz is not None else None,
        "fill": fill,
        "border": {"left": side_of(border.left), "right": side_of(border.right),
                   "top": side_of(border.top), "bottom": side_of(border.bottom)},
        "align": {"h": cell.alignment.horizontal, "v": cell.alignment.vertical,
                  "wrap": bool(cell.alignment.wrap_text)},
    }


def xlsx_model(path: str) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    values: list[list[Any]] = []
    styles: list[list[dict[str, Any]]] = []
    for row_index in range(1, ws.max_row + 1):
        row_values: list[Any] = []
        row_styles: list[dict[str, Any]] = []
        for col_index in range(1, ws.max_column + 1):
            cell = ws.cell(row_index, col_index)
            row_values.append(cell.value)
            row_styles.append(xlsx_cell_sig(cell))
        values.append(row_values)
        styles.append(row_styles)

    registers: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    field_rows = reserved_rows = 0
    for index, row in enumerate(values[1:], start=2):
        if row[3] == "Mixed":
            current = {"name": norm_reg(row[0]), "offset": row[1], "reset": row[4], "fields": []}
            registers.append(current)
        else:
            field_rows += 1
            if row[2] == "RESERVED":
                reserved_rows += 1
            if current is None:
                raise SystemExit(f"field row before first Reg row at sheet row {index}")
            current["fields"].append({"span": row[1], "name": row[2], "access": row[3],
                                      "reset": row[4], "desc": row[5]})

    normalized_values = [list(row) for row in values]
    for row in normalized_values[1:]:
        if row[3] == "Mixed":
            row[0] = norm_reg(row[0])

    return {
        "sheetnames": wb.sheetnames,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "header": values[0],
        "values_normalized": normalized_values,
        "styles": styles,
        "registers": registers,
        "counts": {"register_rows": len(registers), "field_rows": field_rows,
                   "reserved_rows": reserved_rows,
                   "fields_total": sum(len(reg["fields"]) for reg in registers)},
        "row_heights": {str(r): ws.row_dimensions[r].height for r in range(1, ws.max_row + 1)
                        if ws.row_dimensions[r].height is not None},
        "column_widths": {k: v.width for k, v in ws.column_dimensions.items()},
        "merged_cells": sorted(str(m) for m in ws.merged_cells.ranges),
        "freeze_panes": ws.freeze_panes,
    }


def xlsx_style_diffs(ours: dict[str, Any], ecap: dict[str, Any], max_examples: int = 5) -> dict[str, Any]:
    examples: list[dict[str, Any]] = []
    differing = 0
    for r, (row_a, row_b) in enumerate(zip(ours["styles"], ecap["styles"]), start=1):
        for c, (a, b) in enumerate(zip(row_a, row_b), start=1):
            if a != b:
                differing += 1
                if len(examples) < max_examples:
                    examples.append({"row": r, "col": c, "ours": a, "ecap": b})
    return {"agree": differing == 0, "cells_compared": ours["max_row"] * ours["max_column"],
            "differing_cells": differing, "examples": examples}


# ---------------------------------------------------------------------------- docx oracle
def style_map(path: str) -> dict[str, dict[str, Any]]:
    xml = zipfile.ZipFile(path).read("word/styles.xml").decode("utf-8")
    styles: dict[str, dict[str, Any]] = {}
    for match in re.finditer(r'<w:style [^>]*w:styleId="([^"]+)".*?</w:style>', xml, re.S):
        style_id, body = match.group(1), match.group(0)
        based = re.search(r'<w:basedOn w:val="([^"]+)"', body)
        rpr = re.search(r"<w:rPr>(.*?)</w:rPr>", body, re.S)
        bold = None
        if rpr:
            if re.search(r"<w:b\s*/>", rpr.group(1)) or re.search(r'<w:b w:val="(?:1|true)"', rpr.group(1)):
                bold = True
            elif re.search(r'<w:b w:val="(?:0|false)"', rpr.group(1)):
                bold = False
        styles[style_id] = {"based": based.group(1) if based else None, "bold": bold}
    return styles


def effective_bold(styles: dict[str, dict[str, Any]], style_id: str | None) -> bool | None:
    seen: set[str] = set()
    while style_id and style_id not in seen:
        seen.add(style_id)
        style = styles.get(style_id)
        if style is None:
            return None
        if style["bold"] is not None:
            return style["bold"]
        style_id = style["based"]
    return None


def tc_info(tc: Any, styles: dict[str, dict[str, Any]]) -> dict[str, Any]:
    text = "".join(node.text or "" for node in tc.iter(qn("w:t")))
    tc_pr = tc.find(qn("w:tcPr"))
    fill_raw = span = vmerge = width = None
    if tc_pr is not None:
        shd = tc_pr.find(qn("w:shd"))
        fill_raw = None if shd is None else shd.get(qn("w:fill"))
        gs = tc_pr.find(qn("w:gridSpan"))
        span = 1 if gs is None else int(gs.get(qn("w:val")))
        vm = tc_pr.find(qn("w:vMerge"))
        vmerge = None if vm is None else (vm.get(qn("w:val")) or "continue")
        tcw = tc_pr.find(qn("w:tcW"))
        width = None if tcw is None else {"type": tcw.get(qn("w:type")), "w": int(tcw.get(qn("w:w")))}

    run_bolds: list[bool] = []
    for paragraph in tc.findall(qn("w:p")):
        ppr = paragraph.find(qn("w:pPr"))
        pstyle = None
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            pstyle = None if ps is None else ps.get(qn("w:val"))
        for run in paragraph.findall(qn("w:r")):
            rpr = run.find(qn("w:rPr"))
            direct = None
            rstyle = None
            if rpr is not None:
                direct = True if rpr.find(qn("w:b")) is not None else None
                rs = rpr.find(qn("w:rStyle"))
                rstyle = None if rs is None else rs.get(qn("w:val"))
            if direct is not None:
                run_bolds.append(direct)
            elif rstyle is not None and effective_bold(styles, rstyle) is not None:
                run_bolds.append(bool(effective_bold(styles, rstyle)))
            else:
                inherited = effective_bold(styles, pstyle)
                if inherited is not None:
                    run_bolds.append(bool(inherited))
    bold = True if any(run_bolds) else (False if run_bolds else None)
    return {"text": text, "fill_raw": fill_raw,
            "fill": None if fill_raw in (None, "auto") else fill_raw,
            "span": span, "vmerge": vmerge, "bold": bold, "width": width}


def table_model(table: Any, styles: dict[str, dict[str, Any]]) -> dict[str, Any]:
    trs = table._tbl.findall(qn("w:tr"))
    rows = []
    for tr in trs:
        tr_pr = tr.find(qn("w:trPr"))
        height = None
        if tr_pr is not None:
            tr_h = tr_pr.find(qn("w:trHeight"))
            height = None if tr_h is None else tr_h.get(qn("w:val"))
        rows.append({"height": height, "cells": [tc_info(tc, styles) for tc in tr.findall(qn("w:tc"))]})
    grid = [int(g.get(qn("w:w"))) for g in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    return {"rows": len(rows), "cols": len(table.columns), "grid": grid, "trs": rows}


def sect_model(sect_pr: Any) -> dict[str, Any]:
    def attrs(element: Any) -> Any:
        if element is None:
            return None
        return {key.split("}")[-1]: value for key, value in sorted(element.attrib.items())}
    return {"pgSz": attrs(sect_pr.find(qn("w:pgSz"))),
            "pgMar": attrs(sect_pr.find(qn("w:pgMar"))),
            "cols": attrs(sect_pr.find(qn("w:cols"))),
            "docGrid": attrs(sect_pr.find(qn("w:docGrid"))),
            "headerReferences": len(sect_pr.findall(qn("w:headerReference"))),
            "footerReferences": len(sect_pr.findall(qn("w:footerReference")))}


def content_cells(table: dict[str, Any]) -> list[list[dict[str, Any]]]:
    """Cell facts that bear content/visual meaning; width is compared separately."""
    return [[{key: cell[key] for key in ("text", "fill", "span", "vmerge", "bold")}
             for cell in tr["cells"]] for tr in table["trs"]]


def docx_model(path: str) -> dict[str, Any]:
    warnings_seen: list[str] = []

    class Capture(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            warnings_seen.append(record.getMessage())

    root_logger = logging.getLogger()
    handler = Capture(level=logging.WARNING)
    root_logger.addHandler(handler)
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            doc = Document(path)
    finally:
        root_logger.removeHandler(handler)

    styles = style_map(path)
    zip_parts = sorted(zipfile.ZipFile(path).namelist())

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    seq: list[dict[str, Any]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            seq.append({"kind": "p", "style": paragraph.style.name, "text": paragraph.text})
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            seq.append({"kind": "t", "rows": len(table.rows), "cols": len(table.columns)})

    h2_count = sum(1 for item in seq if item["kind"] == "p" and item["style"] == "Heading 2")
    if h2_count == 0:
        raise SystemExit("no Heading 2 in " + path)

    chapter_table_models: list[dict[str, Any]] = []
    seen_h2 = 0
    chapter_started = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            if paragraph.style.name == "Heading 2":
                seen_h2 += 1
                chapter_started = seen_h2 == h2_count
        elif child.tag == qn("w:tbl") and chapter_started:
            chapter_table_models.append(table_model(Table(child, doc), styles))

    chapter_index = next(i for i, item in enumerate(seq)
                         if item["kind"] == "p" and item["style"] == "Heading 2"
                         and sum(1 for prior in seq[:i + 1]
                                 if prior["kind"] == "p" and prior["style"] == "Heading 2") == h2_count)
    chapter_seq = seq[chapter_index:]
    chapter_paragraphs = [item for item in chapter_seq if item["kind"] == "p"]
    chapter_h3 = [item["text"] for item in chapter_paragraphs if item["style"] == "Heading 3"]

    overview_index = next((i for i, table in enumerate(chapter_table_models)
                           if table["rows"] == 63 and table["cols"] == 34), None)
    if overview_index is None:
        raise SystemExit("overview table (63x34) not found in " + path)
    overview = chapter_table_models[overview_index]
    rest = chapter_table_models[overview_index + 1:]
    pairs = [{"bit": rest[i], "detail": rest[i + 1]} for i in range(0, len(rest) - 1, 2)
             if rest[i]["rows"] == 6 and rest[i]["cols"] == 16]

    offset_lines: list[dict[str, str]] = []
    for index, item in enumerate(chapter_paragraphs):
        if item["style"] != "Heading 3" or item["text"].endswith("总览"):
            continue
        offset = reset = ""
        for follower in chapter_paragraphs[index + 1: index + 8]:
            if follower["style"] == "Heading 3":
                break
            if follower["text"].startswith("地址偏移"):
                offset = follower["text"]
            if follower["text"].startswith("复位值"):
                reset = follower["text"]
        offset_lines.append({"offset": offset, "reset": reset})

    def normalize_table(table: dict[str, Any]) -> dict[str, Any]:
        clone = json.loads(json.dumps(table, ensure_ascii=False))
        for tr in clone["trs"]:
            for cell in tr["cells"]:
                cell["text"] = norm_reg(cell["text"])
        return clone

    return {
        "parts": zip_parts,
        "open_warnings": warnings_seen + [str(item.message) for item in caught],
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "h3_total": [item["text"] for item in seq if item["kind"] == "p" and item["style"] == "Heading 3"],
        "chapter_title": seq[chapter_index]["text"],
        "chapter_paragraph_count": len(chapter_paragraphs),
        "chapter_h3": chapter_h3,
        "register_h3_normalized": [norm_reg(title) for title in chapter_h3
                                   if not title.endswith("总览")],
        "chapter_table_count": len(chapter_table_models),
        "overview": normalize_table(overview),
        "register_pairs": [{"bit": normalize_table(pair["bit"]),
                            "detail": normalize_table(pair["detail"])} for pair in pairs],
        "offset_lines": offset_lines,
        "sectPr": sect_model(doc.sections[-1]._sectPr),
        "body_scope": {
            "body_paragraphs_before_chapter": chapter_index,
            "body_tables": len(doc.tables) - len(chapter_table_models),
            "body_h3": len([item for item in seq[:chapter_index]
                            if item["kind"] == "p" and item["style"] == "Heading 3"]),
            "body_table_dims": [{"rows": item["rows"], "cols": item["cols"]}
                                for item in seq[:chapter_index] if item["kind"] == "t"],
        },
    }


# ---------------------------------------------------------------------------- comparison
def width_summary(widths: list[Any]) -> dict[str, Any]:
    types = sorted({width["type"] for width in widths if width is not None})
    return {"types": types, "cells": len(widths), "example": widths[0] if widths else None}


def main() -> None:
    (ours_xlsx, ecap_xlsx, ours_docx_path, ecap_docx_path, run1_path, run2_path,
     verify_path, control_xlsx_path, control_docx_path, outdir) = sys.argv[1:11]
    out = Path(outdir)

    x_ours = xlsx_model(ours_xlsx)
    x_ecap = xlsx_model(ecap_xlsx)
    d_ours = docx_model(ours_docx_path)
    d_ecap = docx_model(ecap_docx_path)

    comparisons: list[dict[str, Any]] = []

    def add(item: str, ours: Any, ecap: Any, note: str = "", expected: bool = False,
            expected_note: str = "") -> dict[str, Any]:
        agree = canon(ours) == canon(ecap)
        entry = {"item": item, "ours": ours, "ecap": ecap, "agree": agree,
                 "note": (expected_note if expected and not agree else note),
                 "classification": ("expected" if expected and not agree
                                    else ("match" if agree else "unexpected"))}
        comparisons.append(entry)
        return entry

    # ---- xlsx
    add("xlsx sheet names", x_ours["sheetnames"], x_ecap["sheetnames"])
    add("xlsx dimensions", x_ours["dimensions"], x_ecap["dimensions"])
    add("xlsx header row (9 texts)", x_ours["header"], x_ecap["header"])
    add("xlsx row model counts", x_ours["counts"], x_ecap["counts"])
    add("xlsx values 162x9 (ECAP_ prefix stripped)",
        {"rows": len(x_ours["values_normalized"]), "cells": x_ours["max_row"] * x_ours["max_column"],
         "digest": digest(x_ours["values_normalized"])},
        {"rows": len(x_ecap["values_normalized"]), "cells": x_ecap["max_row"] * x_ecap["max_column"],
         "digest": digest(x_ecap["values_normalized"])})
    add("xlsx per-register offset/reset/field spans",
        {"registers": len(x_ours["registers"]),
         "fields": sum(len(reg["fields"]) for reg in x_ours["registers"]),
         "digest": digest(x_ours["registers"])},
        {"registers": len(x_ecap["registers"]),
         "fields": sum(len(reg["fields"]) for reg in x_ecap["registers"]),
         "digest": digest(x_ecap["registers"])})
    style_diff = xlsx_style_diffs(x_ours, x_ecap)
    add("xlsx cell styles 162x9 (bold/font/size/fill/border/align)",
        {"digest": digest(x_ours["styles"]), "cells": x_ours["max_row"] * x_ours["max_column"]},
        {"digest": digest(x_ecap["styles"]), "cells": x_ecap["max_row"] * x_ecap["max_column"]},
        note=f"{style_diff['differing_cells']} differing cells; examples={style_diff['examples']}")
    add("xlsx row heights", x_ours["row_heights"], x_ecap["row_heights"])
    add("xlsx column widths", x_ours["column_widths"], x_ecap["column_widths"])
    add("xlsx merged cells", x_ours["merged_cells"], x_ecap["merged_cells"])
    add("xlsx freeze panes", x_ours["freeze_panes"], x_ecap["freeze_panes"])

    # ---- docx package / scope
    ours_parts, ecap_parts = set(d_ours["parts"]), set(d_ecap["parts"])
    add("docx package parts",
        {"count": len(ours_parts), "parts": sorted(ours_parts)},
        {"count": len(ecap_parts), "parts": sorted(ecap_parts)},
        note="ours ships the 9 required OOXML parts; eCap additionally ships "
             + ", ".join(sorted(ecap_parts - ours_parts)))
    add("docx total table count", d_ours["table_count"], d_ecap["table_count"])
    add("docx register-chapter table count", d_ours["chapter_table_count"], d_ecap["chapter_table_count"])
    add("docx register-chapter table sequence",
        {"overview": 1, "bit": len(d_ours["register_pairs"]), "detail": len(d_ours["register_pairs"])},
        {"overview": 1, "bit": len(d_ecap["register_pairs"]), "detail": len(d_ecap["register_pairs"])})
    add("docx register-chapter H2 title", d_ours["chapter_title"], d_ecap["chapter_title"])
    add("docx H3 total (whole document)", len(d_ours["h3_total"]), len(d_ecap["h3_total"]),
        note="scope difference: ours renders the full manual body; eCap is a template clone")
    add("docx register H3 titles (ECAP_ prefix stripped)",
        compact_list(d_ours["register_h3_normalized"]),
        compact_list(d_ecap["register_h3_normalized"]))
    add("docx register-chapter H3 count (incl. overview heading)",
        len(d_ours["chapter_h3"]), len(d_ecap["chapter_h3"]),
        note="eCap emits `IP寄存器总览` H3 before the overview table; ours emits the table directly")
    add("docx register-chapter paragraph count",
        d_ours["chapter_paragraph_count"], d_ecap["chapter_paragraph_count"],
        note="the one-paragraph delta is the same missing overview H3")
    add("docx body scope outside register chapter", d_ours["body_scope"], d_ecap["body_scope"],
        note="eCap's docx is a template clone (版本历史 + placeholder sections); ours renders the real manual body")

    # ---- docx overview
    add("docx overview dims/grid/heights",
        {"rows": d_ours["overview"]["rows"], "cols": d_ours["overview"]["cols"],
         "grid": d_ours["overview"]["grid"],
         "heights": [tr["height"] for tr in d_ours["overview"]["trs"]]},
        {"rows": d_ecap["overview"]["rows"], "cols": d_ecap["overview"]["cols"],
         "grid": d_ecap["overview"]["grid"],
         "heights": [tr["height"] for tr in d_ecap["overview"]["trs"]]})
    add("docx overview header cells (text/span)",
        [[{key: cell[key] for key in ("text", "span")} for cell in tr["cells"]]
         for tr in d_ours["overview"]["trs"][:1]],
        [[{key: cell[key] for key in ("text", "span")} for cell in tr["cells"]]
         for tr in d_ecap["overview"]["trs"][:1]],
        expected=True,
        expected_note="EXPECTED: ours merges the 32 bit columns into one `31..0` cell (span 32); "
                      "eCap prints one cell per bit number (31..0)")
    add("docx overview header cell width type",
        width_summary([cell["width"] for cell in d_ours["overview"]["trs"][0]["cells"]]),
        width_summary([cell["width"] for cell in d_ecap["overview"]["trs"][0]["cells"]]),
        expected=True,
        expected_note="EXPECTED: our overview cell widths are dxa; eCap's template uses pct")
    add("docx overview header bold",
        [cell["bold"] for cell in d_ours["overview"]["trs"][0]["cells"]],
        [cell["bold"] for cell in d_ecap["overview"]["trs"][0]["cells"]],
        note="ours writes direct bold on the header cells; eCap's header is not bold")
    add("docx overview body rows (text/fill/vmerge/span/bold, ECAP_ stripped)",
        summary(content_cells(d_ours["overview"])[1:]),
        summary(content_cells(d_ecap["overview"])[1:]))
    add("docx overview body cell width type",
        width_summary([cell["width"] for tr in d_ours["overview"]["trs"][1:] for cell in tr["cells"]]),
        width_summary([cell["width"] for tr in d_ecap["overview"]["trs"][1:] for cell in tr["cells"]]),
        expected=True,
        expected_note="EXPECTED: our overview cell widths are dxa; eCap's template uses pct")
    add("docx overview cell shd representation (auto vs absent)",
        {"auto_cells": sum(1 for tr in d_ours["overview"]["trs"] for cell in tr["cells"]
                           if cell["fill_raw"] == "auto")},
        {"auto_cells": sum(1 for tr in d_ecap["overview"]["trs"] for cell in tr["cells"]
                           if cell["fill_raw"] == "auto")},
        note="eCap writes <w:shd w:fill=\"auto\"/> on overview cells; ours omits shd. "
             "Both render as no shading; XML representation only.")

    # ---- docx bit/detail tables
    bit_content = [content_cells(pair["bit"]) for pair in d_ours["register_pairs"]]
    ecap_bit_content = [content_cells(pair["bit"]) for pair in d_ecap["register_pairs"]]
    detail_content = [content_cells(pair["detail"]) for pair in d_ours["register_pairs"]]
    ecap_detail_content = [content_cells(pair["detail"]) for pair in d_ecap["register_pairs"]]
    add("docx bit tables content (31 x 6x16: text/fill/span/merge/bold)",
        {"count": len(bit_content), "cells": sum(len(tr) for table in bit_content for tr in table),
         "digest": digest(bit_content)},
        {"count": len(ecap_bit_content),
         "cells": sum(len(tr) for table in ecap_bit_content for tr in table),
         "digest": digest(ecap_bit_content)})
    bit_width_diffs = []
    for index, (ours_pair, ecap_pair) in enumerate(zip(d_ours["register_pairs"], d_ecap["register_pairs"])):
        for row_index, (row_a, row_b) in enumerate(zip(ours_pair["bit"]["trs"], ecap_pair["bit"]["trs"])):
            for cell_index, (cell_a, cell_b) in enumerate(zip(row_a["cells"], row_b["cells"])):
                if cell_a["width"] != cell_b["width"]:
                    bit_width_diffs.append({"table": index, "row": row_index, "cell": cell_index,
                                            "ours": cell_a["width"], "ecap": cell_b["width"]})
    add("docx bit-table cell width (dxa twips)",
        {"differing_cells": len(bit_width_diffs), "examples": bit_width_diffs[:5]},
        {"differing_cells": 0},
        note="ours uses the tblGrid width (519) for every bit cell; eCap's per-cell tcW uses 518 "
             "for 20 of 36 cells (1 twip = 1/1440 inch). tblGrid is identical.")
    add("docx bit-table row heights",
        [tr["height"] for tr in d_ours["register_pairs"][0]["bit"]["trs"]],
        [tr["height"] for tr in d_ecap["register_pairs"][0]["bit"]["trs"]])
    add("docx detail tables content (31: text/fill/span/merge/bold)",
        {"count": len(detail_content), "cells": sum(len(tr) for table in detail_content for tr in table),
         "digest": digest(detail_content)},
        {"count": len(ecap_detail_content),
         "cells": sum(len(tr) for table in ecap_detail_content for tr in table),
         "digest": digest(ecap_detail_content)})
    add("docx detail-table grid (tblGrid, first register)",
        d_ours["register_pairs"][0]["detail"]["grid"], d_ecap["register_pairs"][0]["detail"]["grid"])
    detail_width_types = sorted({cell["width"]["type"] for pair in d_ours["register_pairs"]
                                 for tr in pair["detail"]["trs"] for cell in tr["cells"]})
    ecap_detail_width_types = sorted({cell["width"]["type"] for pair in d_ecap["register_pairs"]
                                      for tr in pair["detail"]["trs"] for cell in tr["cells"]})
    add("docx detail-table cell width type", detail_width_types, ecap_detail_width_types,
        expected=True,
        expected_note="EXPECTED (same class as the overview widths): ours writes dxa "
                      "(830/1659/5807 twips); eCap writes pct")
    add("docx per-register offset/reset lines", compact_list(d_ours["offset_lines"]),
        compact_list(d_ecap["offset_lines"]))
    add("docx sectPr (pgSz/pgMar/cols/docGrid/header+footer refs)", d_ours["sectPr"], d_ecap["sectPr"])
    add("docx python-docx open warnings (ours)", d_ours["open_warnings"], [])

    # ---- determinism
    run1 = json.loads(Path(run1_path).read_text(encoding="utf-8"))
    run2 = json.loads(Path(run2_path).read_text(encoding="utf-8"))
    files_1 = {Path(item["path"]).name: item["sha256"] for item in run1["result"]["files"]}
    files_2 = {Path(item["path"]).name: item["sha256"] for item in run2["result"]["files"]}
    rerun_identical = files_1 == files_2
    add("rerun byte identity (run1 vs run2 sha256)", files_1, files_2,
        note="both runs report the same sha256 for both artifacts")

    # ---- verify tool + controls
    verify = json.loads(Path(verify_path).read_text(encoding="utf-8"))
    control_xlsx = json.loads(Path(control_xlsx_path).read_text(encoding="utf-8"))
    control_docx = json.loads(Path(control_docx_path).read_text(encoding="utf-8"))

    expected_premises = {
        "ECAP_ prefix on core register names": {
            "confirmed": True,
            "normalization": "strip one leading `ECAP_` from register names only; field names untouched",
            "example": "eCap xlsx A2 = ECAP_TSCTR, ours = TSCTR",
        },
        "no cloned 版本历史 table": {
            "confirmed": True,
            "measured": {"ours_total_tables": d_ours["table_count"], "ecap_total_tables": d_ecap["table_count"],
                         "ours_register_chapter_tables": d_ours["chapter_table_count"],
                         "ecap_register_chapter_tables": d_ecap["chapter_table_count"]},
            "note": "the task's `63 vs 64 register-chapter tables` wording does not match measurement: "
                    "both register chapters have 63 tables (1 overview + 31 bit + 31 detail) and both "
                    "documents have 64 tables total. The free slot is outside the chapter: ours is the "
                    "real manual's 信号 table (21x4); eCap's is the cloned 版本历史 table (2x3).",
        },
        "overview header 3 cells vs 32 per-bit numbers": {
            "confirmed": True,
            "ours": "Offset | Register | 31..0 (gridSpan 32, 3 tc)",
            "ecap": "Offset | Register | 31 | 30 | ... | 0 (34 tc)",
        },
        "overview cell widths dxa vs pct": {
            "confirmed": True,
            "extension": "the same dxa vs pct difference also covers every detail-table cell "
                         "(and the overview body), not only the overview header",
        },
    }

    findings: list[dict[str, Any]] = []
    for entry in comparisons:
        if not entry["agree"] and entry["classification"] == "unexpected":
            findings.append({"id": f"F{len(findings) + 1}", "severity": "info",
                             "item": entry["item"], "ours": entry["ours"],
                             "ecap": entry["ecap"], "note": entry["note"]})
    for finding in findings:
        if finding["item"] == "docx register-chapter paragraph count":
            finding["note"] = "same root cause as the missing overview H3 (156 vs 157 paragraphs)"

    wrong_expected = 30
    wrong_count_control = {
        "expected_register_count": wrong_expected,
        "measured_register_count": x_ours["counts"]["register_rows"],
        "control_fires": x_ours["counts"]["register_rows"] != wrong_expected,
        "meaning": "the oracle comparison is falsifiable: asserting 30 registers against a file with 31 "
                   "is reported as a mismatch",
    }
    tampered_xlsx = control_xlsx["result"]
    tampered_docx = control_docx["result"]
    negative_control = {
        "wrong_count": wrong_count_control,
        "tampered_xlsx_header": {
            "artifact": str((out / "negative-control" / "tampered-header-surgical.xlsx").resolve()),
            "mutation": "one byte-string replacement in xl/worksheets/sheet1.xml: "
                        "'register name' -> 'wrong register name'",
            "verify_ok": tampered_xlsx["ok"],
            "failed": tampered_xlsx["failed"],
            "detail": next((check["detail"] for check in tampered_xlsx["checks"] if not check["ok"]), ""),
            "control_fires": tampered_xlsx["ok"] is False and "V1" in tampered_xlsx["failed"],
        },
        "tampered_docx_missing_part": {
            "artifact": str((out / "negative-control" / "tampered-no-settings.docx").resolve()),
            "mutation": "removed word/settings.xml from the OOXML package (1 of the 9 required parts)",
            "verify_ok": tampered_docx["ok"],
            "failed": tampered_docx["failed"],
            "detail": next((check["detail"] for check in tampered_docx["checks"] if not check["ok"]), ""),
            "control_fires": tampered_docx["ok"] is False and "V6" in tampered_docx["failed"],
        },
        "how_a_regressed_format_is_detected": [
            "xlsx value/style drift: the 162x9 value digest or the 162x9 style digest changes, and the "
            "per-register offset/reset/field-span digest changes",
            "docx register drift: the register-H3 digest, overview-body digest, bit-table digest or "
            "detail-table digest changes; table counts and grids are compared numerically",
            "docx package drift: the parts list and sectPr comparison change; V6 requires the 9 parts",
            "determinism drift: rerun sha256 values stop matching",
            "gate drift: mpd_docgen_verify returns ok=false and names the failed V-id",
            "negative controls prove the gate bites: a one-string header edit fails V1 alone; removing "
            "word/settings.xml fails V6 alone",
        ],
    }

    content_items = {
        "xlsx sheet names", "xlsx dimensions", "xlsx header row (9 texts)", "xlsx row model counts",
        "xlsx values 162x9 (ECAP_ prefix stripped)", "xlsx per-register offset/reset/field spans",
        "xlsx cell styles 162x9 (bold/font/size/fill/border/align)", "xlsx row heights",
        "xlsx column widths", "xlsx merged cells", "xlsx freeze panes",
        "docx total table count", "docx register-chapter table count",
        "docx register-chapter table sequence", "docx register H3 titles (ECAP_ prefix stripped)",
        "docx overview dims/grid/heights",
        "docx overview body rows (text/fill/vmerge/span/bold, ECAP_ stripped)",
        "docx bit tables content (31 x 6x16: text/fill/span/merge/bold)",
        "docx bit-table row heights",
        "docx detail tables content (31: text/fill/span/merge/bold)",
        "docx detail-table grid (tblGrid, first register)",
        "docx per-register offset/reset lines",
        "docx sectPr (pgSz/pgMar/cols/docGrid/header+footer refs)",
        "docx python-docx open warnings (ours)",
        "rerun byte identity (run1 vs run2 sha256)",
    }
    content_failures = [entry["item"] for entry in comparisons
                        if entry["item"] in content_items and not entry["agree"]]
    controls_fire = (wrong_count_control["control_fires"]
                     and negative_control["tampered_xlsx_header"]["control_fires"]
                     and negative_control["tampered_docx_missing_part"]["control_fires"])
    content_ok = (not content_failures and verify["result"]["ok"] and rerun_identical and controls_fire)

    generated = [{"name": Path(item["path"]).name, "path": item["path"],
                  "run1_bytes": item["bytes"], "run1_sha256": item["sha256"],
                  "run2_sha256": files_2[Path(item["path"]).name],
                  **({"lang": item["lang"]} if "lang" in item else {})}
                 for item in run1["result"]["files"]]

    verdict = (
        "PASS (content fidelity): "
        f"{x_ours['counts']['register_rows']}/{x_ecap['counts']['register_rows']} registers, "
        f"{x_ours['counts']['field_rows']}/{x_ecap['counts']['field_rows']} field rows, "
        f"{x_ours['counts']['reserved_rows']}/{x_ecap['counts']['reserved_rows']} RESERVED rows; "
        f"{x_ours['max_row'] * x_ours['max_column']}/"
        f"{x_ecap['max_row'] * x_ecap['max_column']} xlsx cells agree in value, and "
        f"{style_diff['cells_compared'] - style_diff['differing_cells']}/{style_diff['cells_compared']} "
        "agree in style; "
        f"{len(d_ours['register_pairs'])}/{len(d_ecap['register_pairs'])} bit tables and "
        f"{len(d_ours['register_pairs'])}/{len(d_ecap['register_pairs'])} detail tables agree in content; "
        f"{d_ours['chapter_table_count']}/{d_ecap['chapter_table_count']} register-chapter tables; "
        f"{len(negative_control['tampered_xlsx_header']['failed'])}-id xlsx and "
        f"{len(negative_control['tampered_docx_missing_part']['failed'])}-id docx negative controls fire; "
        f"verify ok={verify['result']['ok']} ("
        + ",".join(check["id"] for check in verify["result"]["checks"])
        + "). "
        f"{len(findings)} non-expected structural divergence(s) are recorded as findings; "
        "none touches register, field, RESERVED, bit-table or detail-table semantics."
    ) if content_ok else (
        "FAIL (content fidelity): " + ("; ".join(content_failures) if content_failures
                                       else "verify or determinism/controls failed")
        + f"; verify ok={verify['result']['ok']}, rerun_identical={rerun_identical}, "
          f"controls_fire={controls_fire}"
    )

    result = {
        "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "evidence_dir": str(out.resolve()),
        "generator_invoked": {
            "tool": "mpd_docgen_run (packages/mpd-docgen-plugin/src/tool-run.ts, execute())",
            "exec_stub": "{agent:{session:{header:{cwd:<repo root>}}}}",
            "args_run1": run1["args"],
            "args_run2": run2["args"],
            "result_run1": {"ok": run1["result"]["ok"], "registers": run1["result"]["registers"],
                            "fields": run1["result"]["fields"], "warnings": run1["result"]["warnings"]},
        },
        "inputs": [
            {"role": role, "path": str(Path(path).resolve()), "bytes": Path(path).stat().st_size,
             "sha256": sha256(path)}
            for role, path in (("ours_xlsx", ours_xlsx), ("ecap_xlsx", ecap_xlsx),
                               ("ours_docx", ours_docx_path), ("ecap_docx", ecap_docx_path),
                               ("source_manual", SOURCE_MANUAL))
        ],
        "generated": generated,
        "retention": "generated/ and generated-rerun/ are retained: the largest artifact is "
                     f"{max(item['run1_bytes'] for item in generated)} bytes (<100KB each)",
        "rerun": {"byte_identical": rerun_identical, "files": generated,
                  "command": "bun invoke-tools.ts run <dir>/generated-rerun run2"},
        "oracle": {
            "implementation": "pure python: openpyxl 3.1.5 + python-docx 1.2.0; imports no plugin code",
            "xlsx": {
                "ours": {key: x_ours[key] for key in ("sheetnames", "dimensions", "max_row", "max_column",
                                                      "counts", "row_heights", "column_widths",
                                                      "merged_cells", "freeze_panes")},
                "ecap": {key: x_ecap[key] for key in ("sheetnames", "dimensions", "max_row", "max_column",
                                                      "counts", "row_heights", "column_widths",
                                                      "merged_cells", "freeze_panes")},
                "value_digest_ours": digest(x_ours["values_normalized"]),
                "value_digest_ecap": digest(x_ecap["values_normalized"]),
                "style_comparison": style_diff,
            },
            "docx": {
                "ours": {key: d_ours[key] for key in ("parts", "open_warnings", "paragraph_count",
                                                      "table_count", "h3_total", "chapter_title",
                                                      "chapter_paragraph_count", "chapter_h3",
                                                      "register_h3_normalized", "chapter_table_count",
                                                      "sectPr", "body_scope")},
                "ecap": {key: d_ecap[key] for key in ("parts", "open_warnings", "paragraph_count",
                                                      "table_count", "h3_total", "chapter_title",
                                                      "chapter_paragraph_count", "chapter_h3",
                                                      "register_h3_normalized", "chapter_table_count",
                                                      "sectPr", "body_scope")},
                "overview": {
                    "ours": {"rows": d_ours["overview"]["rows"], "cols": d_ours["overview"]["cols"],
                             "grid": d_ours["overview"]["grid"],
                             "heights": [tr["height"] for tr in d_ours["overview"]["trs"]],
                             "header_cells": [[{key: cell[key] for key in ("text", "span", "bold")}
                                               for cell in tr["cells"]]
                                              for tr in d_ours["overview"]["trs"][:1]],
                             "body_digest": digest(content_cells(d_ours["overview"])[1:])},
                    "ecap": {"rows": d_ecap["overview"]["rows"], "cols": d_ecap["overview"]["cols"],
                             "grid": d_ecap["overview"]["grid"],
                             "heights": [tr["height"] for tr in d_ecap["overview"]["trs"]],
                             "header_cells": [[{key: cell[key] for key in ("text", "span", "bold")}
                                               for cell in tr["cells"]]
                                              for tr in d_ecap["overview"]["trs"][:1]],
                             "body_digest": digest(content_cells(d_ecap["overview"])[1:])},
                },
                "bit_tables": {"count": len(d_ours["register_pairs"]), "dims": [6, 16],
                               "tblGrid": d_ours["register_pairs"][0]["bit"]["grid"],
                               "row_heights_ours": [tr["height"] for tr in d_ours["register_pairs"][0]["bit"]["trs"]],
                               "row_heights_ecap": [tr["height"] for tr in d_ecap["register_pairs"][0]["bit"]["trs"]],
                               "content_digest_ours": digest(bit_content),
                               "content_digest_ecap": digest(ecap_bit_content)},
                "detail_tables": {"count": len(d_ours["register_pairs"]),
                                  "tblGrid_ours": d_ours["register_pairs"][0]["detail"]["grid"],
                                  "tblGrid_ecap": d_ecap["register_pairs"][0]["detail"]["grid"],
                                  "width_type_ours": detail_width_types,
                                  "width_type_ecap": ecap_detail_width_types,
                                  "content_digest_ours": digest(detail_content),
                                  "content_digest_ecap": digest(ecap_detail_content)},
            },
        },
        "comparisons": comparisons,
        "expected_premises": expected_premises,
        "verify": {
            "ok": verify["result"]["ok"],
            "checks": verify["result"]["checks"],
            "failed": verify["result"]["failed"],
            "v11": "skipped: V11 requires docxEn; this run generated only the zh-CN docx "
                   "(manualEn intentionally omitted, warning recorded by mpd_docgen_run)",
            "command": "bun invoke-tools.ts verify <ours.xlsx> <ours.zh-CN.docx>",
        },
        "negative_control": negative_control,
        "findings": findings,
        "scope": {
            "written_by_this_lane": [
                str(out.resolve()) + "/invoke-tools.ts",
                str(out.resolve()) + "/oracle.py",
                str(out.resolve()) + "/build-output-log.sh",
                str(out.resolve()) + "/run1.raw", str(out.resolve()) + "/run2.raw",
                str(out.resolve()) + "/verify.raw", str(out.resolve()) + "/sha256sum.txt",
                str(out.resolve()) + "/oracle.stdout", str(out.resolve()) + "/oracle.json",
                str(out.resolve()) + "/result.json", str(out.resolve()) + "/output.log",
                str(out.resolve()) + "/generated/", str(out.resolve()) + "/generated-rerun/",
                str(out.resolve()) + "/negative-control/",
            ],
            "statement": "No source, docs, skills, scripts or package file was edited by this lane; "
                         "every write is inside the evidence directory. `git status` in output.log "
                         "also lists modifications from other lanes (the generator implementation, "
                         "docs, skills) that predate this evidence run.",
        },
        "reproduce": [
            "bun <evidence>/invoke-tools.ts run <evidence>/generated-rerun run2",
            "bun <evidence>/invoke-tools.ts verify <evidence>/generated/eCap_regExcel_full.xlsx "
            "<evidence>/generated/eCap_reg_user_manual.zh-CN.docx",
            "PYTHONPATH=.mpd/tmp/t3/pydeps python3 <evidence>/oracle.py "
            "<ours.xlsx> <ecap.xlsx> <ours.docx> <ecap.docx> <run1.raw> <run2.raw> <verify.raw> "
            "<negative-control/verify-tampered-xlsx.json> <negative-control/verify-tampered-docx.json> "
            "<evidence>",
            "bash <evidence>/build-output-log.sh",
        ],
        "blocking_content_failures": content_failures,
        "verdict": verdict,
        "verdict_rule": "PASS means every content-bearing comparison agrees, the 162x9 xlsx style "
                        "matrix agrees, rerun sha256 values are identical, mpd_docgen_verify returns "
                        "ok=true, and both tampered-artifact controls fail their gate. Recorded "
                        "findings are structural divergences that do not change register/field "
                        "semantics; each is listed with its exact measured values.",
    }

    def docx_summary(model: dict[str, Any]) -> dict[str, Any]:
        return {
            "parts": model["parts"],
            "open_warnings": model["open_warnings"],
            "paragraph_count": model["paragraph_count"],
            "table_count": model["table_count"],
            "h3_total_count": len(model["h3_total"]),
            "chapter_title": model["chapter_title"],
            "chapter_paragraph_count": model["chapter_paragraph_count"],
            "chapter_h3_count": len(model["chapter_h3"]),
            "register_h3_normalized": model["register_h3_normalized"],
            "chapter_table_count": model["chapter_table_count"],
            "overview": {
                "rows": model["overview"]["rows"],
                "cols": model["overview"]["cols"],
                "grid": model["overview"]["grid"],
                "heights": [tr["height"] for tr in model["overview"]["trs"]],
                "header_cells": [[{key: cell[key] for key in ("text", "span", "bold")}
                                  for cell in tr["cells"]] for tr in model["overview"]["trs"][:1]],
                "body_digest": digest(content_cells(model["overview"])[1:]),
            },
            "bit_tables": {"count": len(model["register_pairs"]), "dims": [6, 16],
                           "tblGrid": model["register_pairs"][0]["bit"]["grid"],
                           "row_heights": [tr["height"] for tr in model["register_pairs"][0]["bit"]["trs"]],
                           "content_digest": digest([content_cells(pair["bit"])
                                                     for pair in model["register_pairs"]])},
            "detail_tables": {"count": len(model["register_pairs"]),
                              "tblGrid": model["register_pairs"][0]["detail"]["grid"],
                              "content_digest": digest([content_cells(pair["detail"])
                                                        for pair in model["register_pairs"]])},
            "offset_lines": model["offset_lines"],
            "sectPr": model["sectPr"],
            "body_scope": model["body_scope"],
        }

    (out / "oracle.json").write_text(json.dumps({
        "xlsx": {
            "ours": {key: value for key, value in x_ours.items()
                     if key in ("sheetnames", "dimensions", "max_row", "max_column", "counts",
                                "row_heights", "column_widths", "merged_cells", "freeze_panes")},
            "ecap": {key: value for key, value in x_ecap.items()
                     if key in ("sheetnames", "dimensions", "max_row", "max_column", "counts",
                                "row_heights", "column_widths", "merged_cells", "freeze_panes")},
            "value_digest_ours": digest(x_ours["values_normalized"]),
            "value_digest_ecap": digest(x_ecap["values_normalized"]),
            "style_digest_ours": digest(x_ours["styles"]),
            "style_digest_ecap": digest(x_ecap["styles"]),
            "registers_digest_ours": digest(x_ours["registers"]),
            "registers_digest_ecap": digest(x_ecap["registers"]),
        },
        "docx": {"ours": docx_summary(d_ours), "ecap": docx_summary(d_ecap)},
        "comparisons": comparisons,
        "expected_premises": expected_premises,
        "negative_control": negative_control,
        "findings": findings,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n",
                                     encoding="utf-8")

    print("ORACLE SUMMARY")
    print(f"xlsx counts ours={x_ours['counts']} ecap={x_ecap['counts']}")
    print(f"xlsx styles differing={style_diff['differing_cells']}/{style_diff['cells_compared']}")
    print(f"docx tables ours={d_ours['table_count']} ecap={d_ecap['table_count']}; "
          f"chapter tables {d_ours['chapter_table_count']}/{d_ecap['chapter_table_count']}; "
          f"register pairs {len(d_ours['register_pairs'])}/{len(d_ecap['register_pairs'])}")
    print(f"docx open warnings ours={d_ours['open_warnings']} ecap={d_ecap['open_warnings']}")
    print(f"comparisons: {sum(1 for c in comparisons if c['agree'])} agree, "
          f"{sum(1 for c in comparisons if not c['agree'])} differ, {len(comparisons)} total")
    for entry in comparisons:
        if not entry["agree"]:
            print(f"  {entry['classification'].upper():10s} {entry['item']}")
    print(f"findings={len(findings)} verify_ok={verify['result']['ok']} "
          f"verify_failed={verify['result']['failed']} rerun_identical={rerun_identical} "
          f"controls_fire={controls_fire}")
    print("VERDICT:", verdict)


if __name__ == "__main__":
    main()
